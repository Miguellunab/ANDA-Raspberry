from pathlib import Path
import re

import openpyxl
import requests


# ---------------------------------------------------------
# CONFIGURACIÓN
# ---------------------------------------------------------

CARPETA_SCRIPT = Path(__file__).resolve().parent
ARCHIVO_EXCEL = CARPETA_SCRIPT / "canciones y links.xlsx"
CARPETA_CANCIONES = CARPETA_SCRIPT / "Canciones"


# ---------------------------------------------------------
# FUNCIONES
# ---------------------------------------------------------

def limpiar_nombre(nombre):
    """
    Elimina caracteres que Windows no permite en nombres de archivo.
    """
    nombre = str(nombre).strip()
    nombre = re.sub(r'[<>:"/\\|?*]', "_", nombre)
    return nombre


def descargar_audio(nombre, url):
    nombre = limpiar_nombre(nombre)

    # Agregar .mp3 si no tiene extensión
    if not Path(nombre).suffix:
        nombre += ".mp3"

    destino = CARPETA_CANCIONES / nombre

    # Evitar volver a descargar archivos existentes
    if destino.exists():
        print(f"⏭️  Ya existe: {nombre}")
        return

    try:
        print(f"⬇️  Descargando: {nombre}")

        respuesta = requests.get(
            url,
            stream=True,
            timeout=60
        )

        respuesta.raise_for_status()

        with open(destino, "wb") as archivo:
            for bloque in respuesta.iter_content(chunk_size=1024 * 64):
                if bloque:
                    archivo.write(bloque)

        print(f"✅ Descargado: {nombre}")

    except Exception as e:
        print(f"❌ Error descargando {nombre}: {e}")


# ---------------------------------------------------------
# PROGRAMA PRINCIPAL
# ---------------------------------------------------------

def main():
    if not ARCHIVO_EXCEL.exists():
        print(f"❌ No encontré el archivo:")
        print(ARCHIVO_EXCEL)
        return

    # Crear carpeta Canciones
    CARPETA_CANCIONES.mkdir(exist_ok=True)

    # Abrir Excel
    libro = openpyxl.load_workbook(ARCHIVO_EXCEL, data_only=True)
    hoja = libro.active

    descargadas = 0
    errores = 0

    # En tu archivo los datos empiezan en la fila 4
    for fila in range(4, hoja.max_row + 1):

        nombre = hoja.cell(row=fila, column=2).value  # Columna B
        url = hoja.cell(row=fila, column=3).value     # Columna C

        # Ignorar filas vacías
        if not nombre or not url:
            continue

        if not str(url).startswith(("http://", "https://")):
            print(f"⚠️ Link inválido en fila {fila}: {url}")
            errores += 1
            continue

        try:
            descargar_audio(nombre, str(url).strip())
            descargadas += 1
        except Exception as e:
            print(f"❌ Error en fila {fila}: {e}")
            errores += 1

    print("\n" + "=" * 50)
    print("DESCARGA FINALIZADA")
    print("=" * 50)
    print(f"📁 Carpeta: {CARPETA_CANCIONES}")
    print(f"🎵 Procesadas: {descargadas}")

    if errores:
        print(f"⚠️ Errores: {errores}")


if __name__ == "__main__":
    main()